"""
Génère les données chiffrées de la page de saisie à partir du classeur Excel des notes de CC.

Usage :
    python3 build_data.py <classeur_notes.xlsx> <index.html> <Codes_acces_enseignants.xlsx>

- Chaque onglet (un par enseignant) devient un bloc chiffré (AES-GCM, clé dérivée du code
  d'accès par PBKDF2-SHA256). Sans code, la page publique ne révèle ni noms ni liens.
- Les codes existants (lus dans le fichier de codes s'il existe) sont conservés ; seuls les
  nouveaux enseignants reçoivent un nouveau code.
- Les données sont injectées dans index.html entre les repères DATA-START et DATA-END.
Dépendances : openpyxl, cryptography.
"""
import sys, os, re, json, base64, secrets, hashlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

ITER = 250_000
ALPHA = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"  # sans 0/O, 1/I/L


def clean(v):
    return "" if v is None else str(v).strip()


def parse_header(txt):
    parts = clean(txt).split(" | ")
    out = {"titre": parts[0] if parts else "", "promotion": "", "campus": "", "semestre": "", "code": ""}
    if len(parts) > 1:
        pc = re.split(r"\s+—\s+", parts[1])
        out["promotion"] = pc[0]
        out["campus"] = pc[1] if len(pc) > 1 else ""
    if len(parts) > 2:
        out["semestre"] = parts[2]
    if len(parts) > 3:
        out["code"] = re.sub(r"^Code\s*:\s*", "", parts[3])
    return out


def parse_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    t = {"enseignant": clean(rows[0][0]) or ws.title, "modules": []}
    cur, in_table = None, False
    for r in rows[1:]:
        a = clean(r[0]) if r else ""
        if " | " in a and re.search(r"Code\s*:", a):
            cur = parse_header(a)
            cur["etudiants"] = []
            t["modules"].append(cur)
            in_table = False
            continue
        if cur is None:
            continue
        if a in ("Étudiant", "Etudiant"):
            in_table = True
            continue
        if in_table:
            if not a:
                if not any(clean(x) for x in r):
                    in_table = False
                continue
            if not any(s["nom"] == a for s in cur["etudiants"]):
                cur["etudiants"].append({"nom": a, "notion": clean(r[1]) if len(r) > 1 else ""})
    return t


def new_code(used_ids):
    while True:
        pid = "".join(secrets.choice(ALPHA) for _ in range(4))
        if pid not in used_ids:
            return pid + "-" + "".join(secrets.choice(ALPHA) for _ in range(8))


def encrypt(obj, code):
    salt, iv = secrets.token_bytes(16), secrets.token_bytes(12)
    secret = code.split("-", 1)[1].upper()
    key = hashlib.pbkdf2_hmac("sha256", secret.encode(), salt, ITER, 32)
    ct = AESGCM(key).encrypt(iv, json.dumps(obj, ensure_ascii=False).encode(), None)
    b = lambda x: base64.b64encode(x).decode()
    return {"s": b(salt), "iv": b(iv), "ct": b(ct)}


def main(src, html_path, codes_path):
    wb = openpyxl.load_workbook(src)
    existing = {}
    base_url = "https://VOTRE-COMPTE.github.io/notes-cc/"
    if os.path.exists(codes_path):
        cw = openpyxl.load_workbook(codes_path).active
        base_url = clean(cw["B2"].value) or base_url
        for r in cw.iter_rows(min_row=5, values_only=True):
            if r[0] and r[1]:
                existing[clean(r[0])] = clean(r[1])

    ADMIN = "— Code administrateur —"
    entries, table, catalogue = {}, [], []
    admin_code = existing.get(ADMIN) or new_code(set())
    entries[admin_code.split("-")[0]] = None  # réserve l'identifiant
    for ws in wb.worksheets:
        t = parse_sheet(ws)
        if not t["modules"]:
            continue
        code = existing.get(t["enseignant"]) or new_code(set(entries))
        pid = code.split("-")[0]
        entries[pid] = encrypt(t, code)
        n_stu = sum(len(m["etudiants"]) for m in t["modules"])
        table.append((t["enseignant"], code, len(t["modules"]), n_stu))
        catalogue.append(t)
    entries[admin_code.split("-")[0]] = encrypt({"catalogue": catalogue}, admin_code)

    payload = json.dumps({"iter": ITER, "entries": entries}, separators=(",", ":"))
    html = open(html_path, encoding="utf-8").read()
    html = re.sub(r"/\*DATA-START\*/.*?/\*DATA-END\*/",
                  lambda _: "/*DATA-START*/" + payload + "/*DATA-END*/", html, flags=re.S)
    open(html_path, "w", encoding="utf-8").write(html)

    # Fichier des codes d'accès (à garder confidentiel)
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Codes d'accès"
    teal = PatternFill("solid", fgColor="173A45")
    ws["A1"] = "Codes d'accès à la saisie des notes de CC — document confidentiel, ne pas déposer sur GitHub"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Adresse de la page :"
    ws["B2"] = base_url
    ws["B2"].font = Font(bold=True, color="A8322D")
    ws["A3"] = "Remplacez l'adresse en B2 par celle de votre page GitHub : les liens personnels se mettent à jour."
    ws["A3"].font = Font(italic=True, color="5D6B70")
    heads = ["Enseignant", "Code d'accès", "Modules", "Étudiants", "Lien personnel (à envoyer)"]
    for i, hd in enumerate(heads, 1):
        c = ws.cell(row=4, column=i, value=hd)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = teal
    rows = [(ADMIN, admin_code, sum(t[2] for t in table), sum(t[3] for t in table))] + sorted(table)
    for j, (name, code, nm, ns) in enumerate(rows, 5):
        ws.cell(row=j, column=1, value=name)
        ws.cell(row=j, column=2, value=code).font = Font(name="Consolas")
        ws.cell(row=j, column=3, value=nm)
        ws.cell(row=j, column=4, value=ns)
        if name == ADMIN:
            for c in range(1, 5):
                ws.cell(row=j, column=c).font = Font(bold=True, color="A8322D", name="Consolas" if c == 2 else None)
            ws.cell(row=j, column=5, value="À saisir dans l'onglet « Administration » — ne pas transmettre aux enseignants").font = Font(italic=True, color="A8322D")
            continue
        ws.cell(row=j, column=5, value=f'=HYPERLINK($B$2&"#code="&B{j},$B$2&"#code="&B{j})')
    for col, w in zip("ABCDE", (28, 18, 10, 10, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    out.save(codes_path)
    print(f"{len(table)} enseignants, {sum(t[2] for t in table)} modules, {sum(t[3] for t in table)} lignes étudiants")


if __name__ == "__main__":
    main(*sys.argv[1:4])
